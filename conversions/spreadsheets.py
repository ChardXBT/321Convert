"""Spreadsheet conversions: xlsx -> csv/pdf, csv -> xlsx (openpyxl)."""

import csv
import html
import io
import re
from datetime import date, datetime, time
from pathlib import Path

from core import config
from core.errors import ConversionError
from core.runner import ConversionResult
from core.uploads import read_text

_NUMERIC_VALUE = re.compile(r"^[-+]?\d+([.,]\d+)?$")
_INT_VALUE = re.compile(r"^[-+]?\d+$")
_FLOAT_VALUE = re.compile(r"^[-+]?(\d*\.\d+|\d+\.\d*)$")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def format_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return str(value)


def neutralize(text: str) -> str:
    """Defuse spreadsheet-formula payloads for CSV export."""
    if text.startswith(_FORMULA_PREFIXES) and not _NUMERIC_VALUE.match(text):
        return "'" + text
    return text


def _workbook(path: Path):
    from openpyxl import load_workbook

    try:
        return load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ConversionError("The workbook could not be read. It may be damaged.") from exc


def _sheet(workbook, name):
    if name:
        if name not in workbook.sheetnames:
            raise ConversionError(f"The workbook has no sheet named “{name[:100]}”.")
        return workbook[name]
    return workbook.active


def _iter_sheet(worksheet):
    """All rows as formatted strings, trimmed of trailing empty rows/cols."""
    cells = 0
    rows = []
    width = 0
    for row in worksheet.iter_rows(values_only=True):
        cells += len(row)
        if cells > config.MAX_SPREADSHEET_CELLS:
            raise ConversionError("The spreadsheet contains too many cells.")
        formatted = [format_cell(value) for value in row]
        rows.append(formatted)
        for index, value in enumerate(formatted):
            if value != "":
                width = max(width, index + 1)
    while rows and all(value == "" for value in rows[-1]):
        rows.pop()
    return [[value for value in row[:width]] for row in rows]


def to_csv(input_path: Path, output_path: Path, options: dict) -> ConversionResult:
    workbook = _workbook(input_path)
    try:
        worksheet = _sheet(workbook, options.get("sheet_name"))
        rows = _iter_sheet(worksheet)
    finally:
        workbook.close()
    if not rows:
        raise ConversionError("The selected sheet is empty.")
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        for row in rows:
            writer.writerow([neutralize(value) for value in row])
    return ConversionResult(str(output_path), "csv", [])


def to_pdf(input_path: Path, output_path: Path, options: dict) -> ConversionResult:
    from conversions import pdf_render

    workbook = _workbook(input_path)
    notes = []
    parts = []
    max_columns = 0
    try:
        if options.get("sheet_name"):
            worksheets = [_sheet(workbook, options["sheet_name"])]
        else:
            worksheets = list(workbook.worksheets)
        for worksheet in worksheets:
            rows = _iter_sheet(worksheet)
            if not rows:
                continue
            max_columns = max(max_columns, max(len(row) for row in rows))
            if len(rows) > config.MAX_PDF_TABLE_ROWS:
                notes.append(
                    f"Only the first {config.MAX_PDF_TABLE_ROWS} rows of “{worksheet.title}” were included."
                )
                rows = rows[: config.MAX_PDF_TABLE_ROWS]
            parts.append(_sheet_html(worksheet.title, rows))
    finally:
        workbook.close()
    if not parts:
        raise ConversionError("The workbook is empty.")

    font_size = "8pt" if max_columns > 8 else "9.5pt"
    extra_css = f"table {{ font-size: {font_size}; }}"
    pdf_render.render_pdf(
        "<pdf:nextpage/>".join(parts),
        output_path,
        title=options.get("_stem", ""),
        landscape=max_columns > 6,
        extra_css=extra_css,
    )
    return ConversionResult(str(output_path), "pdf", notes)


def _sheet_html(title: str, rows: list[list[str]]) -> str:
    head = "".join(f"<th>{html.escape(value)}</th>" for value in rows[0])
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(value)}</td>" for value in row) + "</tr>" for row in rows[1:]
    )
    return (
        f"<h2>{html.escape(title)}</h2>"
        f'<table repeat="1"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>'
    )


def _csv_value(text: str):
    if _INT_VALUE.match(text):
        return int(text), False
    if _FLOAT_VALUE.match(text):
        return float(text), False
    return text, True


def csv_to_xlsx(input_path: Path, output_path: Path, options: dict) -> ConversionResult:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell

    text = read_text(input_path)
    try:
        dialect = csv.Sniffer().sniff(text[: 64 * 1024], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(title="Sheet1")
    cells = 0
    for row in csv.reader(io.StringIO(text), dialect):
        cells += len(row)
        if cells > config.MAX_SPREADSHEET_CELLS:
            workbook.close()
            raise ConversionError("The spreadsheet contains too many cells.")
        out = []
        for raw in row:
            value, is_text = _csv_value(raw)
            if is_text:
                # Store literal text even when it starts with "=" so a
                # formula-looking CSV field cannot execute in Excel. openpyxl
                # infers "f" for "=..." on assignment; the data_type override
                # written afterwards is what round-trips as a string.
                cell = WriteOnlyCell(worksheet, value=str(value))
                cell.data_type = "s"
                out.append(cell)
            else:
                out.append(value)
        worksheet.append(out)
    workbook.save(str(output_path))
    return ConversionResult(str(output_path), "xlsx", [])
