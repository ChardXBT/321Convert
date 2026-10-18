"""PDF input conversions: pdf -> txt/docx (pypdf) and pdf -> png/jpg
(pypdfium2 rasterisation)."""

import math
import re
import time
import zipfile
from io import BytesIO
from pathlib import Path

import conversions
from core import config
from core.errors import ConversionError
from core.runner import ConversionResult

_IMAGE_FORMATS = ["JPEG", "PNG", "TIFF", "JPEG2000", "BMP", "GIF", "WEBP"]

_EMPTY_NOTE = "No text could be recognised on these pages."


def _open(path: Path):
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ConversionError("The PDF could not be opened. It may be damaged or corrupted.") from exc
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            raise ConversionError("This PDF is password-protected. Remove the password and try again.")
    try:
        pages = len(reader.pages)
    except Exception as exc:
        raise ConversionError("The PDF could not be opened. It may be damaged or corrupted.") from exc
    if pages == 0:
        raise ConversionError("The PDF has no pages.")
    if pages > config.MAX_PDF_PAGES:
        raise ConversionError(f"PDFs are limited to {config.MAX_PDF_PAGES} pages.")
    return reader


def _page_text(page) -> str:
    try:
        text = page.extract_text() or ""
    except Exception:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n"))


def _ocr_pages(pdf_path: Path, candidates: list[int], notes: list[str]) -> dict[int, str]:
    """Render each candidate page with pypdfium2 and OCR it.

    Bounded by ``config.OCR_MAX_PAGES`` and a wall-clock budget measured from
    the first OCR attempt; pages beyond either limit are simply skipped.
    """
    if not candidates:
        return {}
    import pypdfium2 as pdfium

    from conversions import ocr

    results = {}
    done = 0
    stopped = len(candidates) > config.OCR_MAX_PAGES
    try:
        pdf = pdfium.PdfDocument(str(pdf_path))
    except pdfium.PdfiumError:
        return results
    try:
        started = time.monotonic()
        last_duration = 0.0
        for index in candidates[: config.OCR_MAX_PAGES]:
            elapsed = time.monotonic() - started
            # Predictive: don't start a page that probably would not finish.
            if elapsed > config.OCR_TIME_BUDGET or (
                done and elapsed + last_duration > config.OCR_TIME_BUDGET
            ):
                stopped = True
                break
            page = pdf[index]
            try:
                page_start = time.monotonic()
                image = _render_page(page, scale_cap=config.OCR_RENDER_SCALE, max_side=ocr.MAX_OCR_SIDE)
                try:
                    text = ocr.recognize(image)
                finally:
                    image.close()
                last_duration = time.monotonic() - page_start
            finally:
                page.close()
            done += 1
            if text.strip():
                results[index] = text
    finally:
        pdf.close()
    if stopped:
        notes.append(f"OCR stopped after {done} page(s) to stay within the time limit.")
    if results:
        notes.append(f"Text on {done} page(s) was recognised with OCR; check it for errors.")
    return results


def _scanned_indices(page_texts: list[str]) -> list[int]:
    return [
        index
        for index, text in enumerate(page_texts)
        if len("".join(text.split())) < config.OCR_MIN_TEXT_CHARS
    ]


def to_txt(input_path: Path, output_path: Path, options: dict) -> ConversionResult:
    reader = _open(input_path)
    pages = [_page_text(page) for page in reader.pages]
    notes = []
    for index, text in _ocr_pages(input_path, _scanned_indices(pages), notes).items():
        pages[index] = text
    text = "\n\n".join(pages).strip()
    if not text:
        notes.append(_EMPTY_NOTE)
    output_path.write_text(text, encoding="utf-8")
    return ConversionResult(str(output_path), "txt", notes)


def _paragraphs(text: str) -> list[str]:
    return [" ".join(chunk.split()) for chunk in re.split(r"\n\s*\n", text) if chunk.strip()]


def _pdf_page_images(page, document, budget):
    """Embed up to *budget* images from one page; returns images embedded."""
    Image = conversions._pillow()
    from docx.shared import Inches

    embedded = 0
    try:
        images = list(page.images)
    except Exception:
        return 0
    for image in images:
        if embedded >= budget:
            break
        try:
            data = image.data
        except Exception:  # noqa: S112 -- an unreadable image is skipped
            continue
        if not data or len(data) > config.MAX_EMBEDDED_IMAGE_BYTES:
            continue
        try:
            opened = Image.open(BytesIO(data), formats=_IMAGE_FORMATS)
            opened.load()
            if opened.format == "JPEG":
                payload = data
            else:
                buffer = BytesIO()
                mode = "RGBA" if opened.mode in {"RGBA", "LA", "PA"} else "RGB"
                opened.convert(mode).save(buffer, "PNG")
                payload = buffer.getvalue()
            document.add_picture(BytesIO(payload), width=Inches(6))
            embedded += 1
        except Exception:  # noqa: S112 -- an unusable image is skipped
            continue
    return embedded


def to_docx(input_path: Path, output_path: Path, options: dict) -> ConversionResult:
    from docx import Document

    reader = _open(input_path)
    try:
        document = Document()
    except Exception as exc:
        raise ConversionError("The Word document could not be created.") from exc

    try:
        title = (reader.metadata or {}).get("/Title") if reader.metadata else None
        if isinstance(title, str) and title.strip():
            document.core_properties.title = title.strip()[:200]
    except Exception:  # noqa: S110 -- metadata is best-effort
        pass

    notes = []
    page_texts = [_page_text(page) for page in reader.pages]
    for index, text in _ocr_pages(input_path, _scanned_indices(page_texts), notes).items():
        page_texts[index] = text
    any_text = any(text.strip() for text in page_texts)
    image_budget = config.MAX_EMBEDDED_IMAGES
    for index, page in enumerate(reader.pages):
        text = page_texts[index]
        for paragraph in _paragraphs(text):
            document.add_paragraph(paragraph)
        image_budget -= _pdf_page_images(page, document, image_budget)
        if index < len(reader.pages) - 1:
            document.add_page_break()
    if not any_text:
        notes.append(_EMPTY_NOTE)
    try:
        document.save(str(output_path))
    except Exception as exc:
        raise ConversionError("The Word document could not be created.") from exc
    return ConversionResult(str(output_path), "docx", notes)


def _render_page(page, scale_cap=None, max_side=4000):
    """Render one page, capped at ``max_side`` px on the long side and the
    global image-pixel budget. The returned RGB image owns its bytes."""
    width, height = page.get_size()
    if not width or not height or width * height <= 0:
        raise ConversionError("The PDF has an invalid page size.")
    scale = min(
        scale_cap if scale_cap is not None else config.PDF_RENDER_SCALE,
        max_side / max(width, height),
        math.sqrt(config.MAX_IMAGE_PIXELS / (width * height)),
    )
    bitmap = page.render(scale=scale)
    try:
        pil_image = bitmap.to_pil()
        try:
            return pil_image.convert("RGB")
        finally:
            pil_image.close()
    finally:
        bitmap.close()


def to_images(input_path: Path, output_path: Path, options: dict) -> ConversionResult:
    import pypdfium2 as pdfium

    ext = output_path.suffix.lower().lstrip(".")
    quality = options.get("quality", 82)
    notes = []

    try:
        pdf = pdfium.PdfDocument(str(input_path))
    except pdfium.PdfiumError as exc:
        raise ConversionError(
            "The PDF could not be opened. It may be damaged or password-protected."
        ) from exc
    try:
        total = len(pdf)
        if total == 0:
            raise ConversionError("The PDF has no pages.")
        count = min(total, config.MAX_PDF_RENDER_PAGES)
        if count < total:
            notes.append(f"Only the first {count} pages were rendered.")
        if count == 1:
            page = pdf[0]
            image = _render_page(page)
            if ext == "jpg":
                image.save(output_path, "JPEG", quality=quality)
            else:
                image.save(output_path, "PNG", optimize=False)
            page.close()
            return ConversionResult(str(output_path), ext, notes)

        zip_path = output_path.with_suffix(".zip")
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_STORED) as archive:
            for index in range(count):
                page = pdf[index]
                image = _render_page(page)
                buffer = BytesIO()
                if ext == "jpg":
                    image.save(buffer, "JPEG", quality=quality)
                else:
                    image.save(buffer, "PNG", optimize=False)
                page.close()
                archive.writestr(f"page-{index + 1:03d}.{ext}", buffer.getvalue())
        return ConversionResult(str(zip_path), "zip", notes)
    finally:
        pdf.close()
