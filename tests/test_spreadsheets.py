import io
import json

from openpyxl import load_workbook
from pypdf import PdfReader

from core import config
from tests import samples


def notes_of(response):
    return json.loads(response.headers.get("X-Conversion-Notes", "[]"))


def test_xlsx_to_csv_values(client):
    response = samples.convert(client, "d.xlsx", samples.xlsx(), "csv")
    assert response.status_code == 200
    assert response.mimetype == "text/csv"
    assert response.data[:3] == b"\xef\xbb\xbf"  # BOM for Excel
    text = response.data.decode("utf-8-sig")
    assert "a,3,TRUE,2024-01-05,2024-01-05 13:30:00," in text
    assert "b,2.5,FALSE,,," in text


def test_formula_strings_neutralized(client):
    response = samples.convert(client, "f.xlsx", samples.xlsx_formulas(), "csv")
    assert response.status_code == 200
    text = response.data.decode("utf-8-sig")
    assert "'=cmd" in text
    assert "\n-5" in text or "\r\n-5" in text
    assert "'+1 (555)" in text
    assert "'@handle" in text


def test_sheet_name_selects_second_sheet(client):
    response = samples.convert(client, "two.xlsx", samples.xlsx_two_sheets(), "csv", sheet_name="Second")
    assert response.status_code == 200
    assert b"two" in response.data and b"one" not in response.data


def test_missing_sheet_name_422(client):
    response = samples.convert(client, "two.xlsx", samples.xlsx_two_sheets(), "csv", sheet_name="Nope")
    assert response.status_code == 422
    assert "Nope" in response.data.decode()


def test_empty_sheet_422(client):
    response = samples.convert(client, "e.xlsx", samples.xlsx_empty_sheet(), "csv")
    assert response.status_code == 422


def test_trailing_empty_rows_cols_dropped(client):
    response = samples.convert(client, "d.xlsx", samples.xlsx(), "csv")
    lines = [line for line in response.data.decode("utf-8-sig").splitlines() if line]
    assert len(lines) == 3  # header + 2 data rows, no blank tail


def test_cell_cap(client, monkeypatch):
    monkeypatch.setattr(config, "MAX_SPREADSHEET_CELLS", 10)
    response = samples.convert(client, "d.xlsx", samples.xlsx(), "csv")
    assert response.status_code == 422
    assert b"too many cells" in response.data


def test_xlsx_to_pdf_all_sheets(client):
    response = samples.convert(client, "two.xlsx", samples.xlsx_two_sheets(), "pdf")
    assert response.status_code == 200
    text = "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(response.data)).pages)
    assert "First" in text and "Second" in text
    assert "one" in text and "two" in text


def test_xlsx_to_pdf_named_sheet(client):
    response = samples.convert(client, "two.xlsx", samples.xlsx_two_sheets(), "pdf", sheet_name="Second")
    assert response.status_code == 200
    text = "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(response.data)).pages)
    assert "Second" in text and "First" not in text


def test_xlsx_to_pdf_row_cap(client, monkeypatch):
    monkeypatch.setattr(config, "MAX_PDF_TABLE_ROWS", 1)
    response = samples.convert(client, "d.xlsx", samples.xlsx(), "pdf")
    assert response.status_code == 200
    assert any("first 1 rows" in note for note in notes_of(response))


def test_csv_to_xlsx_sniffs_delimiters(client):
    for delim in (",", ";"):
        response = samples.convert(client, "in.csv", samples.csv_text(delim), "xlsx")
        assert response.status_code == 200, delim
        book = load_workbook(io.BytesIO(response.data))
        sheet = book.active
        assert sheet["A1"].value == "a" and sheet["B1"].value == "b"


def test_csv_to_xlsx_types_and_formula_safety(client):
    response = samples.convert(client, "in.csv", samples.csv_text(","), "xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).active
    assert sheet["A2"].value == 1 and isinstance(sheet["A2"].value, int)
    assert sheet["B2"].value == 2.5 and isinstance(sheet["B2"].value, float)
    formula_cell = sheet["B3"]
    assert formula_cell.value == "=1+1" and formula_cell.data_type == "s"


def test_csv_quoted_newline_survives(client):
    data = io.BytesIO(b'a,b\r\n"line1\nline2",2\r\n')
    response = samples.convert(client, "in.csv", data, "xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).active
    assert sheet["A2"].value == "line1\nline2"


def test_csv_cp1252_decodes(client):
    data = io.BytesIO("café;x".encode("cp1252"))
    response = samples.convert(client, "e.csv", data, "xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).active
    assert sheet["A1"].value == "café"


def test_binary_csv_rejected(client):
    response = samples.convert(client, "b.csv", samples.binary_as_txt(), "xlsx")
    assert response.status_code == 400
    assert b"text file" in response.data
