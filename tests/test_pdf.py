import io
import json
import zipfile

from docx import Document

from core import config
from tests import samples


def notes_of(response):
    return json.loads(response.headers.get("X-Conversion-Notes", "[]"))


def test_pdf_to_txt_two_pages(client):
    response = samples.convert(client, "doc.pdf", samples.pdf_text(["alpha page", "beta page"]), "txt")
    assert response.status_code == 200
    assert response.mimetype == "text/plain"
    text = response.data.decode("utf-8")
    assert "alpha page" in text and "beta page" in text


def test_pdf_to_docx_text_and_page_break(client):
    response = samples.convert(
        client, "doc.pdf", samples.pdf_text(["first page text", "second page text"]), "docx"
    )
    assert response.status_code == 200
    document = Document(io.BytesIO(response.data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "first page text" in text and "second page text" in text
    xml = document.element.xml
    assert "w:br" in xml and 'type="page"' in xml


def test_pdf_to_docx_embeds_images(client):
    response = samples.convert(client, "img.pdf", samples.pdf_with_image(), "docx")
    assert response.status_code == 200
    document = Document(io.BytesIO(response.data))
    assert len(document.inline_shapes) >= 1


def test_owner_only_encrypted_pdf_works(client):
    response = samples.convert(client, "own.pdf", samples.pdf_owner_only(), "txt")
    assert response.status_code == 200
    assert b"secret" in response.data


def test_user_password_pdf_rejected(client):
    response = samples.convert(client, "pw.pdf", samples.pdf_encrypted(), "txt")
    assert response.status_code == 422
    assert b"password-protected" in response.data


def test_blank_pdf_notes_no_text(client):
    response = samples.convert(client, "blank.pdf", samples.pdf_blank(), "txt")
    assert response.status_code == 200
    assert any("No text could be recognised" in note for note in notes_of(response))


def test_pdf_single_page_png(client):
    response = samples.convert(client, "one.pdf", samples.pdf_text(["x"]), "png")
    assert response.status_code == 200
    assert response.mimetype == "image/png"


def test_pdf_render_bitmap_capped(client):
    from PIL import Image

    response = samples.convert(client, "huge.pdf", samples.pdf_huge_page(), "png")
    assert response.status_code == 200
    image = Image.open(io.BytesIO(response.data))
    assert max(image.size) <= 4000


def test_pdf_multi_page_zip(client):
    response = samples.convert(client, "three.pdf", samples.pdf_text(["a", "b", "c"]), "png")
    assert response.status_code == 200
    assert response.mimetype == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(response.data)).namelist()
    assert names == ["page-001.png", "page-002.png", "page-003.png"]


def test_pdf_render_page_cap(client, monkeypatch):
    monkeypatch.setattr(config, "MAX_PDF_RENDER_PAGES", 2)
    response = samples.convert(client, "three.pdf", samples.pdf_text(["a", "b", "c"]), "png")
    assert response.status_code == 200
    assert len(zipfile.ZipFile(io.BytesIO(response.data)).namelist()) == 2
    assert any("first 2 pages" in note for note in notes_of(response))


def test_pdf_to_jpg_quality_changes_size(client):
    small = samples.convert(client, "q.pdf", samples.pdf_text(["text"]), "jpg", quality=20)
    large = samples.convert(client, "q.pdf", samples.pdf_text(["text"]), "jpg", quality=95)
    assert small.status_code == 200 and large.status_code == 200
    assert len(small.data) != len(large.data)


def test_pdf_page_limit(client, monkeypatch):
    monkeypatch.setattr(config, "MAX_PDF_PAGES", 1)
    response = samples.convert(client, "two.pdf", samples.pdf_text(["a", "b"]), "txt")
    assert response.status_code == 422
    assert b"limited to" in response.data


def test_pdf_to_docx_notes_blank(client):
    response = samples.convert(client, "blank.pdf", samples.pdf_blank(), "docx")
    assert response.status_code == 200
    assert any("No text could be recognised" in note for note in notes_of(response))
