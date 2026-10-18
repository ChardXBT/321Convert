import io
import socket

from pypdf import PdfReader

from core import config
from tests import samples


def pdf_text_of(data: bytes) -> str:
    return "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)


def test_docx_to_html_blocks(client):
    response = samples.convert(client, "doc.docx", samples.docx(), "html")
    assert response.status_code == 200
    body = response.data.decode()
    assert "<h1>Main Heading</h1>" in body
    assert "<b>bold words</b>" in body
    assert "<table" in body and "cell-a" in body
    assert '<img src="data:image/png' in body
    assert "<ul>" in body and "<li>item one</li>" in body
    assert "<hr>" in body  # page break maps to hr in HTML output


def test_docx_to_txt(client):
    response = samples.convert(client, "doc.docx", samples.docx(), "txt")
    assert response.status_code == 200
    text = response.data.decode()
    assert "Main Heading" in text
    assert "- item one" in text
    assert "cell-a\tcell-b" in text


def test_docx_to_pdf(client):
    response = samples.convert(client, "doc.docx", samples.docx(), "pdf")
    assert response.status_code == 200
    text = pdf_text_of(response.data)
    assert "Main Heading" in text
    assert "cell-a" in text
    assert len(PdfReader(io.BytesIO(response.data)).pages) >= 2  # page break


def test_damaged_docx_rejected(client):
    response = samples.convert(client, "bad.docx", samples.damaged_docx(), "pdf")
    assert response.status_code == 422


def test_html_to_pdf_no_network_no_scripts(client, monkeypatch):
    def boom(*args, **kwargs):
        raise AssertionError("network!")

    monkeypatch.setattr(socket, "getaddrinfo", boom)
    monkeypatch.setattr(socket, "create_connection", boom)
    payload = (
        "<style>@import url(http://127.0.0.1:1/x.css);"
        "body{background:url(http://127.0.0.1:1/x.png)}</style>"
        "<script>alert(1)</script><iframe src='http://x'></iframe>"
        "<img src='http://127.0.0.1/x'>"
        "<p style='background:url(http://x);color:red'>styled</p>"
    )
    response = samples.convert(client, "page.html", samples.html(payload), "pdf")
    assert response.status_code == 200, response.data[:300]
    assert response.mimetype == "application/pdf"
    assert "styled" in pdf_text_of(response.data)


def test_html_to_pdf_data_image_ok(client):
    pixel = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    payload = f'<img src="data:image/png;base64,{pixel}">'
    response = samples.convert(client, "img.html", samples.html(payload), "pdf")
    assert response.status_code == 200


def test_html_size_limit(client, monkeypatch):
    monkeypatch.setattr(config, "MAX_HTML_BYTES", 100)
    big = samples.html("<p>" + "x" * 500 + "</p>")
    response = samples.convert(client, "big.html", big, "pdf")
    assert response.status_code == 422
    assert b"3 MB" in response.data or b"limited" in response.data


def test_unicode_text_to_pdf_font_embedded(client):
    response = samples.convert(client, "uni.txt", samples.text("Привет мир — Ωmega €"), "pdf")
    assert response.status_code == 200
    assert "Привет мир" in pdf_text_of(response.data)


def test_html_title_becomes_pdf_metadata(client):
    response = samples.convert(client, "t.html", samples.html(), "pdf")
    assert response.status_code == 200
    assert PdfReader(io.BytesIO(response.data)).metadata.title == "Page"


def test_html_title_not_in_pdf_body(client):
    payload = "<html><head><title>Secret Title</title></head><body><h1>Hi</h1></body></html>"
    response = samples.convert(client, "t.html", io.BytesIO(payload.encode()), "pdf")
    assert response.status_code == 200
    reader = PdfReader(io.BytesIO(response.data))
    assert reader.metadata.title == "Secret Title"
    page_text = reader.pages[0].extract_text() or ""
    assert "Secret Title" not in page_text
    assert "Hi" in page_text


def test_markdown_to_html(client):
    response = samples.convert(client, "doc.md", samples.markdown(), "html")
    assert response.status_code == 200
    body = response.data.decode()
    assert "<table>" in body
    assert "<code>" in body or "<pre>" in body
    assert "<title>Heading One</title>" in body


def test_markdown_script_removed(client):
    response = samples.convert(
        client, "doc.md", samples.markdown("# T\n\n<script>alert(1)</script>\n"), "html"
    )
    assert b"<script>" not in response.data


def test_markdown_to_pdf(client):
    response = samples.convert(client, "doc.md", samples.markdown(), "pdf")
    assert response.status_code == 200
    assert "Heading One" in pdf_text_of(response.data)


def test_text_to_html_escapes_and_titles(client):
    response = samples.convert(
        client, "note.txt", samples.text("<script>x</script>"), "html", title="My <b> T"
    )
    assert response.status_code == 200
    body = response.data.decode()
    assert "&lt;script&gt;" in body
    assert "My &lt;b&gt; T" in body
    assert "<title>My &lt;b&gt; T</title>" in body


def test_text_to_html_default_title_is_stem(client):
    response = samples.convert(client, "notes.txt", samples.text("hi"), "html")
    assert "<title>notes</title>" in response.data.decode()


def test_text_to_pdf(client):
    response = samples.convert(client, "n.txt", samples.text("plain words"), "pdf")
    assert response.status_code == 200
    assert "plain words" in pdf_text_of(response.data)


def test_cp1252_text_decodes(client):
    response = samples.convert(client, "c.txt", samples.text_cp1252(), "html")
    assert response.status_code == 200
    assert "Café" in response.data.decode()


def test_cp1252_undefined_byte_replaced(client):
    # 0x81 is undefined in cp1252 and must not fail the decode.
    data = io.BytesIO(b"caf\xe9 \x81 ok")
    response = samples.convert(client, "c.txt", data, "html")
    assert response.status_code == 200
    assert "café" in response.data.decode()


def test_markdown_remote_image_kept_in_html(client):
    source = samples.markdown("# T\n\n![x](https://example.com/a.png)\n")
    response = samples.convert(client, "doc.md", source, "html")
    assert response.status_code == 200
    assert 'src="https://example.com/a.png"' in response.data.decode()


def test_markdown_remote_image_pdf_no_network(client, monkeypatch):
    def boom(*args, **kwargs):
        raise AssertionError("network!")

    monkeypatch.setattr(socket, "getaddrinfo", boom)
    monkeypatch.setattr(socket, "create_connection", boom)
    source = samples.markdown("# T\n\n![x](https://example.com/a.png)\n")
    response = samples.convert(client, "doc.md", source, "pdf")
    assert response.status_code == 200


def test_binary_text_rejected(client):
    response = samples.convert(client, "b.txt", samples.binary_as_txt(), "html")
    assert response.status_code == 400
    assert b"text file" in response.data


# --- DOCX fidelity -----------------------------------------------------------


def _styled_docx() -> io.BytesIO:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("Centered big red")
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    para.add_run("sup").font.superscript = True

    table = document.add_table(rows=2, cols=2)
    cell = table.cell(0, 0)
    cell.width = Inches(2)
    cell.text = "shaded"
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(tc_pr.makeelement(qn("w:shd"), {qn("w:fill"): "FF8800"}))

    image_run = document.add_paragraph().add_run()
    image_run.add_picture(samples.png((8, 8)), width=Inches(2))
    document.add_paragraph("numbered", style="List Number")
    document.sections[0].header.paragraphs[0].text = "HEADER TEXT"
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def test_docx_fidelity_html(client):
    response = samples.convert(client, "s.docx", _styled_docx(), "html")
    assert response.status_code == 200
    body = response.data.decode()
    assert "text-align:center" in body
    assert "font-size:20pt" in body
    assert "color:#ff0000" in body
    assert "background:#ffff00" in body
    assert "<sup>sup</sup>" in body
    assert "background:#ff8800" in body
    assert "width:144pt" in body
    assert 'width="144"' in body
    assert "<ol" in body and "<li" in body
    assert "HEADER TEXT" in body


def test_docx_page_css_landscape():
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm

    from conversions.documents import _page_css

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(297), Mm(210)
    (width, height), _margins = _page_css(document)
    assert 800 < width < 850
    assert 580 < height < 610
    assert width > height


# --- Author stylesheet passthrough -------------------------------------------


def test_sanitize_css_units(monkeypatch):
    from conversions.pdf_render import sanitize_css

    assert "/*" not in sanitize_css("a /* c */ { color: red; }")
    assert "url(x)" not in sanitize_css("@import url(x);p{color:red;}")
    assert "url" not in sanitize_css("@font-face{font-family:X;src:url(x)}p{color:blue;}")
    cleaned = sanitize_css("p{background:url(http://x);color:red;}")
    assert "color:red" in cleaned and "url(" not in cleaned
    cleaned = sanitize_css("p{font-family:Arial;color:blue;}")
    assert "Arial" not in cleaned and "color:blue" in cleaned
    media = sanitize_css("@media print { .a { display:none } }")
    assert "@media" in media and "display:none" in media
    # A declaration ending the block (no ';') is filtered too.
    assert "url(" not in sanitize_css(".a{background:url(http://x)}")
    cleaned = sanitize_css(".a{color:red;background:url(x)}")
    assert "color:red" in cleaned and "url(" not in cleaned
    only = sanitize_css(".a{font-family:Arial}")
    assert "Arial" not in only and ".a{}" in only.replace(" ", "")
    multiline = sanitize_css("p{\ncolor:\nred;\nfont-family:\nArial\n}")
    assert "red" in multiline and "Arial" not in multiline
    monkeypatch.setattr(config, "MAX_CSS_BYTES", 50)
    assert len(sanitize_css("p{color:" + "r" * 500 + ";}")) <= 50


def test_sanitize_html_inline_font_stripped():
    from conversions.pdf_render import sanitize_html

    body = sanitize_html('<p style="font-family:Arial;color:red">x</p>')
    assert "font-family" not in body
    assert "color:red" in body


def test_html_author_css_applies(client):
    payload = '<style>.hide{display:none}</style><p class="hide">SECRETWORD</p><p>Shown</p>'
    response = samples.convert(client, "s.html", samples.html(payload), "pdf")
    assert response.status_code == 200
    text = pdf_text_of(response.data)
    assert "SECRETWORD" not in text
    assert "Shown" in text


# --- Script fonts ------------------------------------------------------------


def test_apply_script_fonts():
    from conversions.pdf_render import apply_script_fonts

    out = apply_script_fonts("<p>Latin 中文 more 안녕 end</p>")
    assert '<span class="x-cjk">中文</span>' in out
    assert '<span class="x-hangul">안녕</span>' in out
    out = apply_script_fonts('<img alt="中文" src="x"><p>&amp; 你好</p>')
    assert 'alt="中文"' in out  # attributes are never wrapped
    assert "&amp;" in out  # entities survive


def test_cjk_txt_to_pdf(client):
    text = "Latin 你好世界 こんにちは 안녕하세요 end"
    response = samples.convert(client, "c.txt", samples.text(text), "pdf")
    assert response.status_code == 200
    extracted = pdf_text_of(response.data)
    for fragment in ("Latin", "你好世界", "こんにちは", "안녕하세요", "end"):
        assert fragment in extracted


def test_cjk_docx_to_pdf(client):
    from docx import Document

    document = Document()
    document.add_paragraph("Latin 你好世界 こんにちは 안녕하세요 end")
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    response = samples.convert(client, "c.docx", buffer, "pdf")
    assert response.status_code == 200
    extracted = pdf_text_of(response.data)
    for fragment in ("你好世界", "こんにちは", "안녕하세요"):
        assert fragment in extracted


def test_cjk_html_to_pdf(client):
    response = samples.convert(
        client, "c.html", samples.html("<p>Latin 你好世界 こんにちは 안녕하세요 end</p>"), "pdf"
    )
    assert response.status_code == 200
    extracted = pdf_text_of(response.data)
    for fragment in ("你好世界", "こんにちは", "안녕하세요"):
        assert fragment in extracted
