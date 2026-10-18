"""Synthetic HTTP probes; optionally enforce the 512 MiB container deployment gate."""

import argparse
import io
import json
import subprocess
import threading
import time
import urllib.error
import urllib.request
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
LINES = [
    "INVOICE TOTAL 1234 USD",
    "Scanned document line two",
    "Another row of text here",
    "Reference number 987654",
    "Final scanned line five",
]
MIB = 1024 * 1024


def request(url, data=None, headers=None):
    req = urllib.request.Request(  # noqa: S310 -- explicit operator URL
        url, data=data, headers=headers or {}
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as response:  # noqa: S310 -- explicit operator URL
            return response.status, response.read()
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read()


def convert(base, filename, source, target, expected=200):
    boundary = uuid.uuid4().hex
    body = (
        (
            f'--{boundary}\r\nContent-Disposition: form-data; name="target"\r\n\r\n{target}\r\n'
            f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n'
            "Content-Type: application/octet-stream\r\n\r\n"
        ).encode()
        + source
        + f"\r\n--{boundary}--\r\n".encode()
    )
    started = time.monotonic()
    status, result = request(
        base + "/api/convert", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"}
    )
    print(f"{filename} -> {target}: HTTP {status}, {time.monotonic() - started:.2f}s", flush=True)
    if status != expected:
        raise RuntimeError(
            f"{filename} -> {target}: expected {expected}, received {status}: {result[:300]!r}"
        )
    return result


def encoded(image, fmt, **kwargs):
    output = io.BytesIO()
    image.save(output, fmt, **kwargs)
    return output.getvalue()


def check(condition, message):
    if not condition:
        raise RuntimeError(message)


def check_ocr(data, pages=1):
    text = " ".join(data.decode("utf-8").upper().split())
    for line in LINES:
        check(text.count(line.upper()) == pages, f"OCR output missing/duplicating synthetic line: {line}")


def docker(container, *args):
    return subprocess.check_output(  # noqa: S603
        ["docker", *args[:1], container, *args[1:]],  # noqa: S607
        text=True,
    ).strip()


def gunicorn_identity(container):
    rows = docker(container, "top", "-eo", "pid,args").splitlines()
    return sorted(row.split()[0] for row in rows if "gunicorn" in row and "multiprocessing" not in row)


def container_memory(container):
    return int(docker(container, "exec", "cat", "/sys/fs/cgroup/memory.peak"))


def smoke(base, container=None):
    status, health = request(base + "/health")
    check(status == 200 and json.loads(health)["status"] == "ok", "Initial health check failed")
    print(f"health: {health.decode().strip()}", flush=True)
    identity = gunicorn_identity(container) if container else None
    if container:
        check(len(identity) == 2, "Expected one Gunicorn master and one worker")
        check(
            docker(container, "exec", "cat", "/sys/fs/cgroup/memory.max") == str(512 * MIB),
            "Container is not limited to 512 MiB",
        )

    stop = threading.Event()
    health_errors = []

    def monitor():
        while not stop.wait(0.5):
            try:
                code, body = request(base + "/health")
                if code != 200 or json.loads(body) != json.loads(health):
                    health_errors.append(f"health returned {code}, or version changed")
            except Exception as exc:
                health_errors.append(type(exc).__name__)

    thread = threading.Thread(target=monitor, daemon=True)
    thread.start()
    try:
        small = encoded(Image.new("RGB", (64, 48), "navy"), "PNG")
        jpg = convert(base, "sample.png", small, "jpg")
        check(Image.open(io.BytesIO(jpg)).format == "JPEG", "Invalid JPEG")
        heic = convert(base, "sample.png", small, "heic")
        jpg = convert(base, "roundtrip.heic", heic, "jpg")
        check(Image.open(io.BytesIO(jpg)).size == (64, 48), "HEIC roundtrip failed")

        text = "Hello document world\n中文 日本語 한국어\n"
        pdf = convert(base, "sample.txt", text.encode(), "pdf")
        extracted = "".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
        check(
            all(value in extracted for value in ("Hello", "中文", "日本語", "한국어")),
            "PDF text/font loss",
        )

        document = Document()
        document.add_heading("Conversion smoke", 1)
        document.add_paragraph("Word document content")
        document.add_table(rows=1, cols=2).cell(0, 0).text = "Table cell"
        buffer = io.BytesIO()
        document.save(buffer)
        pdf = convert(base, "sample.docx", buffer.getvalue(), "pdf")
        check("Table cell" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text(), "DOCX table lost")
        pdf = convert(base, "sample.html", b"<style>p {color: blue}</style><p>HTML content</p>", "pdf")
        check("HTML content" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text(), "HTML content lost")
        pdf = convert(base, "sample.md", b"# Markdown content\n\nA paragraph.", "pdf")
        check("Markdown content" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text(), "Markdown lost")

        workbook = Workbook()
        workbook.active.append(["Label", "Value"])
        workbook.active.append(["Example", 1234])
        buffer = io.BytesIO()
        workbook.save(buffer)
        csv = convert(base, "sample.xlsx", buffer.getvalue(), "csv")
        check("Example,1234" in csv.decode("utf-8-sig"), "Spreadsheet values lost")
        xlsx = convert(base, "sample.csv", csv, "xlsx")
        check(load_workbook(io.BytesIO(xlsx)).active["B2"].value == 1234, "CSV roundtrip failed")
        pdf = convert(base, "sample.xlsx", buffer.getvalue(), "pdf")
        check("Example" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text(), "Spreadsheet PDF lost")
        convert(base, "bad.exe", b"MZ not an executable", "pdf", expected=400)

        image = Image.new("RGB", (1240, 1754), "white")
        font = ImageFont.truetype(str(ROOT / "conversions/fonts/DejaVuSans.ttf"), 34)
        draw = ImageDraw.Draw(image)
        for index, line in enumerate(LINES):
            draw.text((80, 120 + index * 80), line, fill="black", font=font)
        scan_png = encoded(image, "PNG")
        scan_pdf = encoded(image, "PDF", save_all=True, append_images=[image])
        check_ocr(convert(base, "scan.png", scan_png, "txt"))
        check_ocr(convert(base, "scan.pdf", scan_pdf, "txt"), pages=2)
        docx = convert(base, "scan.pdf", scan_pdf, "docx")
        document = Document(io.BytesIO(docx))
        check_ocr("\n".join(p.text for p in document.paragraphs).encode(), pages=2)

        if container:
            # Fresh children can get different native allocator layouts.
            # Repeat the previously intermittent PDF path without retrying or
            # suppressing failures; every request must succeed on its first try.
            for _ in range(6):
                check_ocr(convert(base, "repeat.pdf", scan_pdf, "txt"), pages=2)

        # Two OCR requests and an ordinary conversion arriving together must not
        # overcommit the free instance. The application serializes this work.
        with ThreadPoolExecutor(max_workers=3) as pool:
            first = pool.submit(convert, base, "parallel.png", scan_png, "txt")
            second = pool.submit(convert, base, "parallel.pdf", scan_pdf, "txt")
            third = pool.submit(convert, base, "parallel.txt", b"Concurrent document", "pdf")
            check_ocr(first.result())
            check_ocr(second.result(), pages=2)
            check(bool(PdfReader(io.BytesIO(third.result())).pages), "Concurrent PDF failed")

        # Exercise memory rejection only inside the disposable container.
        # A 36 MP image passes dimension validation but needs large decoded copies.
        if container:
            large = encoded(Image.new("RGB", (6000, 6000), "navy"), "PNG")
            convert(base, "large.png", large, "jpg", expected=422)
        convert(base, "after.png", small, "jpg")
    finally:
        stop.set()
        thread.join(timeout=125)
    check(not health_errors, f"Health failures during conversions: {health_errors}")
    status, after = request(base + "/health")
    check(status == 200 and json.loads(after) == json.loads(health), "Service changed after probes")
    if container:
        check(gunicorn_identity(container) == identity, "Gunicorn restarted during the probes")
        events = docker(container, "exec", "cat", "/sys/fs/cgroup/memory.events")
        event_counts = dict(line.split() for line in events.splitlines())
        check(int(event_counts["oom"]) == 0 and int(event_counts["oom_kill"]) == 0, "Container hit OOM")
        peak = container_memory(container)
        print(f"container peak memory: {peak / MIB:.1f} MiB; limit: 512 MiB; oom events: 0", flush=True)
        check(peak < 448 * MIB, "Less than 64 MiB of memory headroom remains")
    print("All service probes passed; health stayed stable.", flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--container", help="Docker container to check for resource regressions")
    args = parser.parse_args()
    smoke(args.base_url.rstrip("/"), args.container)
